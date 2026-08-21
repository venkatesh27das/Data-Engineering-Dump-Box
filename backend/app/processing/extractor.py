import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from app.processing.layout_detector import (
    RegionCandidate,
    combined_headers,
    detect_implicit_regions,
)
from app.processing.ooxml_inspector import inspect_ooxml
from app.processing.vba_inspector import inspect_vba

FORMULA_REF = re.compile(r"(?:'([^']+)'|([A-Za-z0-9_ ]+))!\$?([A-Z]{1,3})\$?(\d+)")


def workbook_protection_present(book: Any) -> bool:
    security = getattr(book, "security", None)
    return bool(
        security
        and (
            getattr(security, "lockStructure", False)
            or getattr(security, "lockWindows", False)
        )
    )


@dataclass
class ExtractionResult:
    manifest: dict[str, Any]
    sheets: list[dict[str, Any]] = field(default_factory=list)
    regions: list[dict[str, Any]] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    columns: list[dict[str, Any]] = field(default_factory=list)
    formulas: list[dict[str, Any]] = field(default_factory=list)
    named_ranges: list[dict[str, Any]] = field(default_factory=list)
    images: list[dict[str, Any]] = field(default_factory=list)
    charts: list[dict[str, Any]] = field(default_factory=list)
    comments: list[dict[str, Any]] = field(default_factory=list)
    external_links: list[dict[str, Any]] = field(default_factory=list)
    connections: list[dict[str, Any]] = field(default_factory=list)
    queries: list[dict[str, Any]] = field(default_factory=list)
    pivots: list[dict[str, Any]] = field(default_factory=list)
    conditional_formats: list[dict[str, Any]] = field(default_factory=list)
    data_validations: list[dict[str, Any]] = field(default_factory=list)
    macros: list[dict[str, Any]] = field(default_factory=list)
    forms: list[dict[str, Any]] = field(default_factory=list)
    units: list[dict[str, Any]] = field(default_factory=list)
    nodes: list[dict[str, Any]] = field(default_factory=list)
    edges: list[dict[str, Any]] = field(default_factory=list)
    reviews: list[dict[str, Any]] = field(default_factory=list)
    assets: list[dict[str, Any]] = field(default_factory=list)
    table_rows: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    media: list[tuple[str, bytes]] = field(default_factory=list)


def extract_workbook(
    path: Path,
    workbook_id: str,
    run_id: str,
    model_available: bool,
    directives: list[dict[str, Any]] | None = None,
) -> ExtractionResult:
    directives = directives or []
    extension = path.suffix.lower()
    if extension == ".xlsb":
        return _extract_xlsb(path, workbook_id, run_id, model_available, directives)
    keep_vba = extension == ".xlsm"
    book = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_vba=keep_vba, keep_links=True)
    package_features = inspect_ooxml(path, book)
    vba_report = inspect_vba(path) if package_features["macro_present"] else None
    manifest: dict[str, Any] = {
        "package_version": "0.1.0",
        "workbook_id": workbook_id,
        "run_id": run_id,
        "source_file": path.name,
        "file_type": extension.removeprefix("."),
        "macro_present": bool(package_features["macro_present"]),
        "macro_inspection": vba_report or {"inspection_status": "not_present"},
        "protection_present": workbook_protection_present(book),
        "semantic_mode": "local_model" if model_available else "deterministic_fallback",
        "calculation_notice": "Formula expressions were inspected; values were not recalculated.",
        "sheet_count": len(book.sheetnames),
        "sheets": [],
    }
    result = ExtractionResult(manifest=manifest)
    _apply_workbook_features(result, package_features, vba_report)
    workbook_node = _node(f"workbook.{workbook_id}", "workbook", path.name, {"source_file": path.name})
    result.nodes.append(workbook_node)

    for defined in book.defined_names.values():
        item = {
            "id": f"named.{uuid4().hex[:10]}",
            "name": defined.name,
            "value": defined.attr_text,
            "hidden": bool(defined.hidden),
        }
        result.named_ranges.append(item)
        result.assets.append(
            _asset(item["id"], "named_range", defined.name, defined.attr_text or "", None, None, 0.96)
        )

    for sheet_index, sheet in enumerate(book.worksheets, start=1):
        if any(
            item.get("type") == "exclude_sheet"
            and str(item.get("target", {}).get("sheet_name", "")).casefold() == sheet.title.casefold()
            for item in directives
        ):
            continue
        sheet_id = f"{workbook_id}.sheet_{sheet_index:02d}"
        used_range = sheet.calculate_dimension()
        sheet_info = {
            "sheet_id": sheet_id,
            "name": sheet.title,
            "index": sheet_index,
            "visibility": sheet.sheet_state,
            "used_range": used_range,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
            "formula_count": sum(1 for row in sheet.iter_rows() for cell in row if cell.data_type == "f"),
            "table_count": len(sheet.tables),
            "image_count": len(getattr(sheet, "_images", [])),
            "chart_count": len(getattr(sheet, "_charts", [])),
        }
        manifest["sheets"].append(sheet_info)
        result.sheets.append(sheet_info)
        result.assets.append(
            _asset(
                sheet_id,
                "sheet",
                sheet.title,
                f"{sheet.sheet_state} sheet using {used_range}",
                sheet.title,
                used_range,
                0.99,
            )
        )
        result.nodes.append(_node(sheet_id, "sheet", sheet.title, {"source_range": used_range}))
        result.edges.append(
            _edge(
                f"edge.{uuid4().hex[:12]}",
                f"workbook.{workbook_id}",
                "CONTAINS",
                sheet_id,
                ["workbook_structure"],
                1.0,
                sheet.title,
                used_range,
            )
        )
        _extract_sheet_tables(result, sheet, sheet_id, workbook_id, path.name, directives)
        _extract_formulas(result, sheet, sheet_id)
        _extract_comments(result, sheet, sheet_id)
        _extract_visuals(result, sheet, sheet_id)

    for unit in result.units:
        result.assets.append(
            _asset(
                unit["unit_id"],
                "semantic_unit",
                unit["title"],
                unit["text_content"][:240],
                unit["structural_context"].get("sheet_name"),
                unit["provenance"].get("source_range"),
                unit["quality"]["semantic_confidence"],
            )
        )
    for edge in result.edges:
        result.assets.append(
            _asset(
                edge["edge_id"],
                "relationship",
                edge["relationship_type"].replace("_", " ").title(),
                edge["relationship_description"],
                edge["provenance"].get("source_sheet"),
                edge["provenance"].get("source_range"),
                edge["confidence"],
            )
        )

    manifest["counts"] = {
        "sheets": len(result.sheets),
        "tables": len(result.tables),
        "formulas": len(result.formulas),
        "images": len(result.images),
        "charts": len(result.charts),
        "comments": len(result.comments),
        "external_links": len(result.external_links),
        "connections": len(result.connections),
        "queries": len(result.queries),
        "pivots": len(result.pivots),
        "conditional_formats": len(result.conditional_formats),
        "data_validations": len(result.data_validations),
        "macro_modules": len(result.macros),
        "forms": len(result.forms),
        "semantic_units": len(result.units),
        "graph_nodes": len(result.nodes),
        "graph_edges": len(result.edges),
    }
    manifest["complexity"] = _complexity(result)
    if not model_available:
        manifest["warnings"] = [
            "LM Studio was unavailable; semantic assets use deterministic contextual text."
        ]
    _validate(result)
    return result


def _apply_workbook_features(
    result: ExtractionResult,
    features: dict[str, Any],
    vba_report: dict[str, Any] | None,
) -> None:
    result.external_links.extend(features.get("external_links", []))
    result.connections.extend(features.get("connections", []))
    result.queries.extend(features.get("queries", []))
    result.pivots.extend(features.get("pivots", []))
    result.conditional_formats.extend(features.get("conditional_formats", []))
    result.data_validations.extend(features.get("data_validations", []))
    if vba_report:
        result.macros.extend(vba_report.get("modules", []))
    collections = [
        ("external_link", result.external_links, "External link", 0.92),
        ("connection", result.connections, "Workbook connection", 0.91),
        ("query", result.queries, "Workbook query", 0.88),
        ("pivot_table", result.pivots, "Pivot table", 0.9),
        ("conditional_format", result.conditional_formats, "Conditional formatting", 0.96),
        ("data_validation", result.data_validations, "Data validation", 0.97),
        ("macro_module", result.macros, "VBA module", 0.95),
    ]
    for asset_type, items, label, confidence in collections:
        for index, item in enumerate(items, start=1):
            item_id = item.get("pivot_id") or f"{asset_type}.{uuid4().hex[:12]}"
            title = (
                item.get("name") or item.get("module_name") or item.get("source_range") or f"{label} {index}"
            )
            sheet = item.get("sheet_name")
            source_range = item.get("source_range")
            result.assets.append(
                _asset(
                    item_id,
                    asset_type,
                    str(title),
                    _feature_summary(label, item),
                    sheet,
                    source_range,
                    confidence,
                )
            )
            if asset_type == "pivot_table":
                result.nodes.append(_node(item_id, "pivot_table", str(title), {"source_range": source_range}))
    if result.external_links:
        result.reviews.append(
            _review(
                None,
                "external_dependency",
                "Review external workbook dependencies",
                f"Detected {len(result.external_links)} external relationship(s).",
                json.dumps(result.external_links[:10], default=str),
                "Confirm that linked sources are expected and still available.",
                0.82,
                "medium",
            )
        )
    indicators = (vba_report or {}).get("indicators", [])
    risky = [item for item in indicators if item.get("severity") == "high"]
    if risky:
        result.reviews.append(
            _review(
                None,
                "macro_security_indicator",
                "Review static VBA security indicators",
                f"Static inspection found {len(risky)} high-priority VBA indicator(s).",
                json.dumps(risky[:12], default=str),
                "Review the macro source outside this application; macros were not executed.",
                0.95,
                "high",
            )
        )


def _feature_summary(label: str, item: dict[str, Any]) -> str:
    fields = []
    for key in (
        "relationship_type",
        "target",
        "description",
        "type",
        "query_type",
        "source_sheet",
        "source_data_range",
        "rule_type",
        "validation_type",
        "line_count",
    ):
        if item.get(key) not in (None, ""):
            fields.append(f"{key.replace('_', ' ')}: {item[key]}")
    return f"{label}. " + ("; ".join(fields[:6]) or "Metadata extracted from the workbook package.")


def _extract_sheet_tables(
    result: ExtractionResult,
    sheet: Any,
    sheet_id: str,
    workbook_id: str,
    source_file: str,
    directives: list[dict[str, Any]],
) -> None:
    candidates = [
        RegionCandidate(
            name=table.name,
            ref=table.ref,
            region_type="table",
            header_row=range_boundaries(table.ref)[1],
            header_depth=1,
            confidence=0.98,
            detection_method="excel_table",
        )
        for table in sheet.tables.values()
    ]
    if not candidates:
        candidates = detect_implicit_regions(sheet)
    for table_index, candidate in enumerate(candidates, start=1):
        name, ref = candidate.name, candidate.ref
        min_col, min_row, max_col, max_row = range_boundaries(ref)
        if candidate.region_type == "form":
            _extract_form_region(
                result,
                sheet,
                sheet_id,
                workbook_id,
                source_file,
                candidate,
                table_index,
            )
            continue
        header_row = candidate.header_row or min_row
        header_depth = candidate.header_depth or 1
        override = next(
            (
                item
                for item in directives
                if item.get("type") == "override_header_row"
                and str(item.get("target", {}).get("sheet_name", "*")).casefold()
                in {"*", sheet.title.casefold()}
            ),
            None,
        )
        if override:
            requested = int(override.get("parameters", {}).get("header_row", header_row))
            if min_row <= requested <= max_row:
                header_row = requested
                header_depth = max(1, int(override.get("parameters", {}).get("header_depth", 1)))
        header_values, combined = combined_headers(sheet, header_row, header_depth, min_col, max_col)
        headers = _unique_headers(combined)
        header_confidence = candidate.confidence
        data_start_row = min(max_row + 1, header_row + header_depth)
        table_id = f"{sheet_id}.table_{table_index:02d}"
        table = {
            "table_id": table_id,
            "sheet_id": sheet_id,
            "name": name,
            "source_range": ref,
            "header_row": header_row,
            "header_rows": list(range(header_row, header_row + header_depth)),
            "header_depth": header_depth,
            "row_count": max(0, max_row - data_start_row + 1),
            "column_count": len(headers),
            "original_headers": header_values,
            "normalized_headers": headers,
            "confidence": round(min(0.99, 0.72 + 0.26 * header_confidence), 3),
            "region_type": candidate.region_type,
            "detection_method": candidate.detection_method,
        }
        result.tables.append(table)
        result.regions.append(
            {
                "region_id": f"{sheet_id}.region_{table_index:02d}",
                "sheet_id": sheet_id,
                "region_type": candidate.region_type,
                "source_range": ref,
                "confidence": table["confidence"],
                "detection_method": candidate.detection_method,
            }
        )
        result.assets.append(
            _asset(
                table_id,
                "table",
                name,
                f"{table['row_count']} rows and {table['column_count']} columns",
                sheet.title,
                ref,
                table["confidence"],
            )
        )
        result.nodes.append(
            _node(table_id, "table", name, {"source_sheet": sheet.title, "source_range": ref})
        )
        result.edges.append(
            _edge(
                f"edge.{uuid4().hex[:12]}",
                sheet_id,
                "CONTAINS",
                table_id,
                ["table_range"],
                1.0,
                sheet.title,
                ref,
            )
        )
        rows: list[dict[str, Any]] = []
        for col_index, (original, normalized) in enumerate(
            zip(header_values, headers, strict=False), start=min_col
        ):
            column_id = f"{table_id}.column_{col_index - min_col + 1:02d}"
            result.columns.append(
                {
                    "column_id": column_id,
                    "table_id": table_id,
                    "original_name": str(original or ""),
                    "normalized_name": normalized,
                    "source_column": get_column_letter(col_index),
                }
            )
        for row_number in range(data_start_row, min(max_row, data_start_row + 4999) + 1):
            values = [_value(sheet.cell(row_number, column)) for column in range(min_col, max_col + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            record = dict(zip(headers, values, strict=False))
            record["_source_row"] = row_number
            rows.append(record)
            unit_id = f"{table_id}.row_{row_number:05d}"
            labeled = "; ".join(
                f"{header}: {value}"
                for header, value in zip(headers, values, strict=False)
                if value not in (None, "")
            )
            source_row_range = (
                f"{get_column_letter(min_col)}{row_number}:{get_column_letter(max_col)}{row_number}"
            )
            result.units.append(
                {
                    "unit_id": unit_id,
                    "unit_type": "table_record",
                    "title": f"{name} record {row_number}",
                    "text_content": labeled,
                    "structured_content": {
                        key: value for key, value in record.items() if not key.startswith("_")
                    },
                    "semantic_context": {
                        "domain": "Workbook data",
                        "entity_type": name,
                        "business_terms": headers[:12],
                    },
                    "structural_context": {
                        "workbook_id": workbook_id,
                        "sheet_id": sheet_id,
                        "sheet_name": sheet.title,
                        "table_id": table_id,
                        "row_number": row_number,
                    },
                    "media_references": [],
                    "relationships": [],
                    "provenance": {
                        "source_file": source_file,
                        "source_range": source_row_range,
                        "source_cells": [
                            f"{get_column_letter(column)}{row_number}"
                            for column in range(min_col, max_col + 1)
                        ],
                        "extraction_method": "table_parser",
                        "processor_version": "0.1.0",
                    },
                    "quality": {
                        "extraction_confidence": table["confidence"],
                        "semantic_confidence": 0.72,
                        "validation_status": "passed",
                    },
                    "embedding_status": "ready",
                    "entity_extraction_status": "ready",
                }
            )
        result.table_rows[table_id] = rows
        if header_confidence < 0.75:
            result.reviews.append(
                _review(
                    table_id,
                    "uncertain_header",
                    f"Confirm the header for {name}",
                    f"The detected header beginning on row {header_row} is ambiguous.",
                    f"{sheet.title}!{ref}",
                    "Select the correct header row.",
                    0.62,
                    "medium",
                )
            )


def _extract_formulas(result: ExtractionResult, sheet: Any, sheet_id: str) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type != "f":
                continue
            formula_id = f"{sheet_id}.formula_{cell.coordinate}"
            expression = str(cell.value)
            references = []
            for match in FORMULA_REF.finditer(expression):
                target_sheet = (match.group(1) or match.group(2) or "").strip()
                target_cell = f"{match.group(3)}{match.group(4)}"
                references.append({"sheet": target_sheet, "cell": target_cell})
                result.edges.append(
                    _edge(
                        f"edge.{uuid4().hex[:12]}",
                        f"{target_sheet}!{target_cell}",
                        "FEEDS",
                        f"{sheet.title}!{cell.coordinate}",
                        ["formula_reference"],
                        0.98,
                        sheet.title,
                        cell.coordinate,
                    )
                )
            business_rule = (
                f"Calculates {sheet.title}!{cell.coordinate} using "
                f"{len(references)} cross-sheet reference(s)."
            )
            item = {
                "formula_id": formula_id,
                "sheet_id": sheet_id,
                "sheet_name": sheet.title,
                "cell": cell.coordinate,
                "original_expression": expression,
                "normalized_expression": expression.replace("$", ""),
                "business_rule": business_rule,
                "inputs": references,
                "output": f"{sheet.title}!{cell.coordinate}",
                "category": _formula_category(expression),
            }
            result.formulas.append(item)
            result.assets.append(
                _asset(
                    formula_id,
                    "formula",
                    f"Formula {sheet.title}!{cell.coordinate}",
                    item["business_rule"],
                    sheet.title,
                    cell.coordinate,
                    0.9,
                )
            )


def _extract_form_region(
    result: ExtractionResult,
    sheet: Any,
    sheet_id: str,
    workbook_id: str,
    source_file: str,
    candidate: RegionCandidate,
    region_index: int,
) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(candidate.ref)
    title_value = sheet.cell(min_row, min_col).value
    title_row = (
        sum(sheet.cell(min_row, column).value not in (None, "") for column in range(min_col, max_col + 1))
        == 1
    )
    form_id = f"{sheet_id}.form_{region_index:02d}"
    fields: list[dict[str, Any]] = []
    for row_number in range(min_row + (1 if title_row else 0), max_row + 1):
        label = _value(sheet.cell(row_number, min_col))
        values = [
            _value(sheet.cell(row_number, column))
            for column in range(min_col + 1, max_col + 1)
            if _value(sheet.cell(row_number, column)) not in (None, "")
        ]
        if label in (None, "") or not values:
            continue
        source_range = f"{get_column_letter(min_col)}{row_number}:{get_column_letter(max_col)}{row_number}"
        field = {
            "label": str(label),
            "value": values[0] if len(values) == 1 else values,
            "source_range": source_range,
        }
        fields.append(field)
        normalized_label = _unique_headers([label])[0]
        unit_id = f"{form_id}.field_{row_number:04d}"
        result.units.append(
            {
                "unit_id": unit_id,
                "unit_type": "form_field",
                "title": f"{candidate.name}: {label}",
                "text_content": f"{label}: {field['value']}",
                "structured_content": {normalized_label: field["value"]},
                "semantic_context": {
                    "domain": "Workbook form",
                    "entity_type": str(title_value or candidate.name),
                    "business_terms": [str(label)],
                },
                "structural_context": {
                    "workbook_id": workbook_id,
                    "sheet_id": sheet_id,
                    "sheet_name": sheet.title,
                    "form_id": form_id,
                    "row_number": row_number,
                },
                "media_references": [],
                "relationships": [],
                "provenance": {
                    "source_file": source_file,
                    "source_range": source_range,
                    "source_cells": [
                        f"{get_column_letter(column)}{row_number}" for column in range(min_col, max_col + 1)
                    ],
                    "extraction_method": "label_value_form_parser",
                    "processor_version": "0.2.0",
                },
                "quality": {
                    "extraction_confidence": candidate.confidence,
                    "semantic_confidence": 0.78,
                    "validation_status": "passed",
                },
                "embedding_status": "ready",
                "entity_extraction_status": "ready",
            }
        )
    form = {
        "form_id": form_id,
        "sheet_id": sheet_id,
        "sheet_name": sheet.title,
        "title": str(title_value or candidate.name),
        "source_range": candidate.ref,
        "field_count": len(fields),
        "fields": fields,
        "confidence": candidate.confidence,
        "detection_method": candidate.detection_method,
    }
    result.forms.append(form)
    result.regions.append(
        {
            "region_id": form_id,
            "sheet_id": sheet_id,
            "region_type": "form",
            "source_range": candidate.ref,
            "confidence": candidate.confidence,
            "detection_method": candidate.detection_method,
        }
    )
    result.assets.append(
        _asset(
            form_id,
            "form",
            form["title"],
            f"Label-value form with {len(fields)} fields",
            sheet.title,
            candidate.ref,
            candidate.confidence,
        )
    )
    result.nodes.append(_node(form_id, "form", form["title"], {"source_range": candidate.ref}))
    result.edges.append(
        _edge(
            f"edge.{uuid4().hex[:12]}",
            sheet_id,
            "CONTAINS",
            form_id,
            ["label_value_layout"],
            candidate.confidence,
            sheet.title,
            candidate.ref,
        )
    )


def _extract_comments(result: ExtractionResult, sheet: Any, sheet_id: str) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if not cell.comment:
                continue
            item_id = f"{sheet_id}.comment_{cell.coordinate}"
            item = {
                "comment_id": item_id,
                "sheet_name": sheet.title,
                "cell": cell.coordinate,
                "author": cell.comment.author,
                "text": cell.comment.text,
            }
            result.comments.append(item)
            result.assets.append(
                _asset(
                    item_id,
                    "comment",
                    f"Comment at {cell.coordinate}",
                    cell.comment.text[:240],
                    sheet.title,
                    cell.coordinate,
                    0.99,
                )
            )


def _extract_visuals(result: ExtractionResult, sheet: Any, sheet_id: str) -> None:
    for index, image in enumerate(getattr(sheet, "_images", []), start=1):
        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        cell = f"{get_column_letter(marker.col + 1)}{marker.row + 1}" if marker else None
        image_id = f"{sheet_id}.image_{index:03d}"
        extension = getattr(image, "format", "png") or "png"
        try:
            binary = image._data()
            filename = f"{image_id}.{extension}"
            result.media.append((filename, binary))
        except Exception:
            filename = ""
        item = {
            "image_id": image_id,
            "sheet_name": sheet.title,
            "anchor_cell": cell,
            "width": image.width,
            "height": image.height,
            "classification": "embedded_image",
            "media_file": filename,
            "description": (
                "Embedded workbook image; local vision interpretation was not required "
                "for structural extraction."
            ),
            "confidence": 0.76,
        }
        result.images.append(item)
        result.assets.append(
            _asset(image_id, "image", f"Image on {sheet.title}", item["description"], sheet.title, cell, 0.76)
        )
        if not cell:
            result.reviews.append(
                _review(
                    image_id,
                    "uncertain_visual_anchor",
                    "Confirm image context",
                    "The image anchor could not be resolved.",
                    sheet.title,
                    "Associate the image with a table or record.",
                    0.55,
                    "low",
                )
            )
    for index, chart in enumerate(getattr(sheet, "_charts", []), start=1):
        chart_id = f"{sheet_id}.chart_{index:03d}"
        title = _chart_title(chart) or f"Chart {index}"
        series = _chart_series(chart)
        anchor_cell = _drawing_anchor(chart)
        data_sources = sorted(
            {
                formula
                for item in series
                for formula in (item.get("categories_formula"), item.get("values_formula"))
                if formula
            }
        )
        item = {
            "chart_id": chart_id,
            "sheet_name": sheet.title,
            "title": title,
            "chart_type": type(chart).__name__,
            "anchor_cell": anchor_cell,
            "series_count": len(series),
            "series": series,
            "data_sources": data_sources,
            "x_axis_title": _axis_title(getattr(chart, "x_axis", None)),
            "y_axis_title": _axis_title(getattr(chart, "y_axis", None)),
            "style": getattr(chart, "style", None),
            "legend_position": getattr(getattr(chart, "legend", None), "position", None),
            "summary": (
                f"{type(chart).__name__} titled '{title}' with {len(series)} data series"
                + (f" sourced from {', '.join(data_sources[:3])}" if data_sources else "")
                + "."
            ),
            "interpretation_method": "chart_metadata",
            "confidence": 0.88 if data_sources else 0.78,
        }
        result.charts.append(item)
        result.assets.append(
            _asset(
                chart_id,
                "chart",
                title,
                item["summary"],
                sheet.title,
                anchor_cell,
                item["confidence"],
            )
        )


def _extract_xlsb(
    path: Path,
    workbook_id: str,
    run_id: str,
    model_available: bool,
    directives: list[dict[str, Any]],
) -> ExtractionResult:
    from pyxlsb import open_workbook

    vba_report = inspect_vba(path)
    manifest = {
        "package_version": "0.1.0",
        "workbook_id": workbook_id,
        "run_id": run_id,
        "source_file": path.name,
        "file_type": "xlsb",
        "macro_present": bool(vba_report.get("contains_macros")),
        "macro_inspection": vba_report,
        "semantic_mode": "local_model" if model_available else "deterministic_fallback",
        "warnings": [
            "XLSB values and tabular regions were extracted; formula expressions, visibility, and "
            "drawing metadata are limited by the binary parser."
        ],
        "sheets": [],
    }
    result = ExtractionResult(manifest=manifest)
    result.macros.extend(vba_report.get("modules", []))
    result.nodes.append(_node(f"workbook.{workbook_id}", "workbook", path.name, {"source_file": path.name}))
    memory_book = openpyxl.Workbook()
    first_sheet = True
    with open_workbook(path) as book:
        for index, name in enumerate(book.sheets, start=1):
            if any(
                item.get("type") == "exclude_sheet"
                and str(item.get("target", {}).get("sheet_name", "")).casefold() == name.casefold()
                for item in directives
            ):
                continue
            with book.get_sheet(name) as sheet:
                values = [[cell.v for cell in row] for _, row in zip(range(5001), sheet.rows(), strict=False)]
            memory_sheet = memory_book.active if first_sheet else memory_book.create_sheet()
            first_sheet = False
            memory_sheet.title = name
            for row in values:
                memory_sheet.append(row)
            sheet_id = f"{workbook_id}.sheet_{index:02d}"
            max_columns = max((len(row) for row in values), default=0)
            info = {
                "sheet_id": sheet_id,
                "name": name,
                "index": index,
                "visibility": "visible",
                "used_range": f"A1:{get_column_letter(max_columns or 1)}{len(values) or 1}",
                "max_row": len(values),
                "max_column": max_columns,
            }
            result.sheets.append(info)
            manifest["sheets"].append(info)
            result.assets.append(
                _asset(
                    sheet_id,
                    "sheet",
                    name,
                    "Best-effort XLSB sheet extraction",
                    name,
                    info["used_range"],
                    0.72,
                )
            )
            result.nodes.append(_node(sheet_id, "sheet", name, {"source_range": info["used_range"]}))
            result.edges.append(
                _edge(
                    f"edge.{uuid4().hex[:12]}",
                    f"workbook.{workbook_id}",
                    "CONTAINS",
                    sheet_id,
                    ["xlsb_workbook_structure"],
                    0.9,
                    name,
                    info["used_range"],
                )
            )
            _extract_sheet_tables(result, memory_sheet, sheet_id, workbook_id, path.name, directives)
    for unit in result.units:
        result.assets.append(
            _asset(
                unit["unit_id"],
                "semantic_unit",
                unit["title"],
                unit["text_content"][:240],
                unit["structural_context"].get("sheet_name"),
                unit["provenance"].get("source_range"),
                unit["quality"]["semantic_confidence"],
            )
        )
    for edge in result.edges:
        result.assets.append(
            _asset(
                edge["edge_id"],
                "relationship",
                edge["relationship_type"].replace("_", " ").title(),
                edge["relationship_description"],
                edge["provenance"].get("source_sheet"),
                edge["provenance"].get("source_range"),
                edge["confidence"],
            )
        )
    manifest["counts"] = {
        "sheets": len(result.sheets),
        "tables": len(result.tables),
        "forms": len(result.forms),
        "semantic_units": len(result.units),
        "graph_nodes": len(result.nodes),
        "graph_edges": len(result.edges),
        "macro_modules": len(result.macros),
    }
    manifest["complexity"] = _complexity(result)
    result.reviews.append(
        _review(
            result.sheets[0]["sheet_id"] if result.sheets else None,
            "unsupported_workbook_feature",
            "Review XLSB extraction",
            "Formula expressions, hidden-sheet state, and drawing metadata may be incomplete.",
            path.name,
            "Confirm important output structures.",
            0.55,
            "medium",
        )
    )
    return result


def _value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, str) and value.startswith("="):
        return value
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _unique_headers(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for index, value in enumerate(values, start=1):
        base = (
            re.sub(r"[^a-z0-9]+", "_", str(value or f"column_{index}").strip().lower()).strip("_")
            or f"column_{index}"
        )
        seen[base] = seen.get(base, 0) + 1
        result.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return result


def _asset(
    asset_id: str,
    asset_type: str,
    title: str,
    summary: str,
    sheet: str | None,
    source_range: str | None,
    confidence: float,
) -> dict[str, Any]:
    return {
        "id": asset_id,
        "asset_type": asset_type,
        "title": title,
        "summary": summary,
        "source_uri": f"workbook://{sheet or 'workbook'}/{source_range or ''}",
        "content_uri": "",
        "source_sheet": sheet,
        "source_range": source_range,
        "confidence": confidence,
        "review_status": "open" if confidence < 0.7 else "not_required",
    }


def _node(node_id: str, node_type: str, label: str, provenance: dict[str, Any]) -> dict[str, Any]:
    return {"node_id": node_id, "node_type": node_type, "label": label, "provenance": provenance}


def _edge(
    edge_id: str,
    source: str,
    relation: str,
    target: str,
    evidence: list[str],
    confidence: float,
    sheet: str | None,
    source_range: str | None,
) -> dict[str, Any]:
    return {
        "edge_id": edge_id,
        "source_id": source,
        "relationship_type": relation,
        "target_id": target,
        "relationship_description": f"{source} {relation.lower().replace('_', ' ')} {target}",
        "evidence": evidence,
        "confidence": confidence,
        "provenance": {"source_sheet": sheet, "source_range": source_range},
    }


def _review(
    asset_id: str | None,
    review_type: str,
    title: str,
    description: str,
    evidence: str,
    action: str,
    confidence: float,
    severity: str,
) -> dict[str, Any]:
    return {
        "id": f"review.{uuid4().hex[:12]}",
        "asset_id": asset_id,
        "review_type": review_type,
        "title": title,
        "description": description,
        "evidence": evidence,
        "suggested_action": action,
        "confidence": confidence,
        "severity": severity,
    }


def _formula_category(expression: str) -> str:
    upper = expression.upper()
    for name in ("LOOKUP", "SUM", "IF", "AVERAGE", "COUNT", "DATE"):
        if name in upper:
            return name.lower()
    return "calculation"


def _chart_title(chart: Any) -> str | None:
    try:
        paragraphs = chart.title.tx.rich.p
        return " ".join(run.t for paragraph in paragraphs for run in paragraph.r if run.t)
    except Exception:
        return None


def _axis_title(axis: Any) -> str | None:
    if axis is None:
        return None
    try:
        paragraphs = axis.title.tx.rich.p
        return " ".join(run.t for paragraph in paragraphs for run in paragraph.r if run.t)
    except Exception:
        return None


def _drawing_anchor(drawing: Any) -> str | None:
    marker = getattr(getattr(drawing, "anchor", None), "_from", None)
    if marker is None:
        return None
    return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"


def _chart_series(chart: Any) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for index, series in enumerate(getattr(chart, "ser", []), start=1):
        title = None
        series_title = getattr(series, "tx", None)
        if series_title is not None:
            title = getattr(series_title, "v", None)
            title = title or _nested_attr(series_title, "strRef", "f")
        categories = (
            _nested_attr(series, "cat", "strRef", "f")
            or _nested_attr(series, "cat", "numRef", "f")
            or _nested_attr(series, "xVal", "numRef", "f")
        )
        values = _nested_attr(series, "val", "numRef", "f") or _nested_attr(series, "yVal", "numRef", "f")
        items.append(
            {
                "index": index,
                "title": title or f"Series {index}",
                "categories_formula": categories,
                "values_formula": values,
            }
        )
    return items


def _nested_attr(value: Any, *attributes: str) -> Any:
    current = value
    for attribute in attributes:
        current = getattr(current, attribute, None)
        if current is None:
            return None
    return current


def _complexity(result: ExtractionResult) -> dict[str, Any]:
    return {
        "layout": min(1.0, (len(result.regions) + len(result.images)) / 12),
        "data": min(1.0, sum(item["row_count"] for item in result.tables) / 10_000),
        "dependency": min(1.0, len(result.edges) / 500),
        "computation": min(1.0, len(result.formulas) / 1000),
        "visual": min(1.0, (len(result.images) + len(result.charts)) / 20),
        "archetype": "analytical_model"
        if result.formulas
        else "multi-table data package"
        if len(result.tables) > 1
        else "flat dataset",
    }


def _validate(result: ExtractionResult) -> None:
    unit_without_provenance = [unit["unit_id"] for unit in result.units if not unit.get("provenance")]
    if unit_without_provenance:
        result.reviews.append(
            _review(
                None,
                "missing_provenance",
                "Some units lack provenance",
                f"{len(unit_without_provenance)} units need source locations.",
                json.dumps(unit_without_provenance[:10]),
                "Re-run deterministic extraction.",
                0.2,
                "high",
            )
        )
