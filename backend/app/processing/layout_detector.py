from dataclasses import dataclass
from typing import Any

from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class RegionCandidate:
    name: str
    ref: str
    region_type: str
    header_row: int | None
    header_depth: int
    confidence: float
    detection_method: str


def detect_implicit_regions(sheet: Any) -> list[RegionCandidate]:
    """Detect separated tables, repeated blocks, and label/value forms."""
    populated = {
        (cell.row, cell.column) for row in sheet.iter_rows() for cell in row if cell.value not in (None, "")
    }
    if not populated:
        return []
    row_bands = _bands(sorted({row for row, _column in populated}))
    bounds: list[tuple[int, int, int, int]] = []
    for min_row, max_row in row_bands:
        columns = sorted({column for row, column in populated if min_row <= row <= max_row})
        for min_col, max_col in _bands(columns):
            cells = [
                (row, column)
                for row, column in populated
                if min_row <= row <= max_row and min_col <= column <= max_col
            ]
            if cells:
                bounds.append(
                    (
                        min(row for row, _column in cells),
                        min(column for _row, column in cells),
                        max(row for row, _column in cells),
                        max(column for _row, column in cells),
                    )
                )
    candidates: list[RegionCandidate] = []
    repeated = len(bounds) > 1
    for index, (min_row, min_col, max_row, max_col) in enumerate(bounds, start=1):
        if max_row == min_row:
            continue
        ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        if _looks_like_form(sheet, min_row, min_col, max_row, max_col):
            candidates.append(
                RegionCandidate(
                    name=f"{sheet.title} Form {index}",
                    ref=ref,
                    region_type="form",
                    header_row=None,
                    header_depth=0,
                    confidence=0.88,
                    detection_method="label_value_layout",
                )
            )
            continue
        header_row, header_depth, confidence = detect_header(sheet, min_row, min_col, max_row, max_col)
        candidates.append(
            RegionCandidate(
                name=(f"{sheet.title} Block {index}" if repeated else f"{sheet.title} Data"),
                ref=ref,
                region_type="repeated_table" if repeated else "table",
                header_row=header_row,
                header_depth=header_depth,
                confidence=confidence,
                detection_method=("repeated_block_detector" if repeated else "contiguous_region_detector"),
            )
        )
    return candidates


def detect_header(
    sheet: Any, min_row: int, min_col: int, max_row: int, max_col: int
) -> tuple[int, int, float]:
    width = max_col - min_col + 1
    search_end = min(max_row, min_row + 4)
    header_row = min_row
    for row_number in range(min_row, search_end + 1):
        nonempty = sum(
            sheet.cell(row_number, column).value not in (None, "") for column in range(min_col, max_col + 1)
        )
        if nonempty >= max(2, width // 2):
            header_row = row_number
            break
    first = _expanded_row(sheet, header_row, min_col, max_col)
    first_density = sum(value not in (None, "") for value in first) / max(1, width)
    first_strings = _string_ratio(first)
    depth = 1
    if header_row < max_row:
        second = _expanded_row(sheet, header_row + 1, min_col, max_col)
        second_strings = _string_ratio(second)
        merged = any(
            merged.min_row <= header_row <= merged.max_row
            and merged.max_col >= min_col
            and merged.min_col <= max_col
            for merged in sheet.merged_cells.ranges
        )
        if second_strings >= 0.6 and (merged or first_density < 0.75):
            depth = 2
    confidence = 0.68 + 0.2 * first_density + 0.1 * first_strings
    if depth == 2:
        confidence += 0.02
    return header_row, depth, round(min(0.99, confidence), 3)


def combined_headers(
    sheet: Any, header_row: int, header_depth: int, min_col: int, max_col: int
) -> tuple[list[Any], list[str]]:
    rows = [
        _expanded_row(sheet, row, min_col, max_col) for row in range(header_row, header_row + header_depth)
    ]
    originals: list[Any] = []
    combined: list[str] = []
    for offset, column in enumerate(range(min_col, max_col + 1), start=1):
        parts: list[str] = []
        for row in rows:
            value = row[column - min_col]
            text = str(value).strip() if value not in (None, "") else ""
            if text and (not parts or parts[-1].casefold() != text.casefold()):
                parts.append(text)
        original = " ".join(parts) or f"Column {offset}"
        originals.append(original)
        combined.append(original)
    return originals, combined


def _expanded_row(sheet: Any, row: int, min_col: int, max_col: int) -> list[Any]:
    values = [sheet.cell(row, column).value for column in range(min_col, max_col + 1)]
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row:
            source = sheet.cell(merged.min_row, merged.min_col).value
            for column in range(max(min_col, merged.min_col), min(max_col, merged.max_col) + 1):
                values[column - min_col] = source
    return values


def _looks_like_form(sheet: Any, min_row: int, min_col: int, max_row: int, max_col: int) -> bool:
    width = max_col - min_col + 1
    height = max_row - min_row + 1
    if width > 4 or height < 3:
        return False
    row_counts = [
        sum(sheet.cell(row, column).value not in (None, "") for column in range(min_col, max_col + 1))
        for row in range(min_row, max_row + 1)
    ]
    title_row = row_counts[0] == 1
    label_value_rows = 0
    for row in range(min_row + (1 if title_row else 0), max_row + 1):
        label = sheet.cell(row, min_col).value
        value_cells = [sheet.cell(row, column).value for column in range(min_col + 1, max_col + 1)]
        if isinstance(label, str) and label.strip() and any(value not in (None, "") for value in value_cells):
            label_value_rows += 1
    evaluated = height - (1 if title_row else 0)
    return title_row and evaluated >= 2 and label_value_rows / max(1, evaluated) >= 0.7


def _bands(values: list[int]) -> list[tuple[int, int]]:
    if not values:
        return []
    result: list[tuple[int, int]] = []
    start = previous = values[0]
    for value in values[1:]:
        if value > previous + 1:
            result.append((start, previous))
            start = value
        previous = value
    result.append((start, previous))
    return result


def _string_ratio(values: list[Any]) -> float:
    populated = [value for value in values if value not in (None, "")]
    return sum(isinstance(value, str) for value in populated) / max(1, len(populated))
