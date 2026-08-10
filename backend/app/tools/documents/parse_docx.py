from pathlib import Path

from docx import Document

from app.domain.normalized import DocumentPage, NormalizedSource
from app.tools.documents.chunk_document import chunk_document


class DocxParseError(ValueError):
    pass


def parse_docx(path: Path, *, source_id: str, source_name: str) -> NormalizedSource:
    try:
        document = Document(path)
    except Exception as error:
        raise DocxParseError(f"DOCX could not be opened: {error}") from error
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            blocks.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = "\n".join(blocks)
    page = DocumentPage(page_number=1, text=text, character_count=len(text))
    confidence = 1.0 if len(text) >= 100 else max(0.25, len(text) / 100)
    warnings = [] if text else ["DOCX did not contain extractable text"]
    return NormalizedSource(
        source_id=source_id,
        source_name=source_name,
        modality="unstructured",
        source_type="docx",
        text=text,
        pages=[page],
        chunks=chunk_document([page]),
        parsing_confidence=confidence,
        warnings=warnings,
    )
