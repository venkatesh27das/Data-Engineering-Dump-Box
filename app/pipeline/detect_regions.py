"""Whitespace segmentation with explicit table/merged-title precedence and explainable scores."""

from collections import defaultdict
from collections.abc import Iterable
from hashlib import sha256

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from app.models.region import Region
from app.models.workbook import Sheet


def bounds(points: Iterable[tuple[int, int]]) -> tuple[int, int, int, int]:
    points = list(points)
    return (
        min(c for r, c in points),
        min(r for r, c in points),
        max(c for r, c in points),
        max(r for r, c in points),
    )


def ref_for(box: tuple[int, int, int, int]) -> str:
    left, top, right, bottom = box
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def contains(box: tuple[int, int, int, int], row: int, col: int) -> bool:
    return box[0] <= col <= box[2] and box[1] <= row <= box[3]


def split_blocks(points: set[tuple[int, int]]) -> list[set[tuple[int, int]]]:
    """Split only on entirely empty rows/columns; blanks inside a table survive."""
    pending, result = [points], []
    while pending:
        block = pending.pop()
        if not block:
            continue
        for axis in (0, 1):
            values = sorted({p[axis] for p in block})
            group_for, group, previous = {}, 0, values[0]
            for value in values:
                if value > previous + 1:
                    group += 1
                group_for[value] = group
                previous = value
            if group:
                groups = defaultdict(set)
                for point in block:
                    groups[group_for[point[axis]]].add(point)
                pending.extend(groups.values())
                break
        else:
            result.append(block)
    return sorted(result, key=lambda block: (bounds(block)[1], bounds(block)[0]))


def detect_regions(
    ws: Worksheet, sheet: Sheet, run_id: str, threshold: float = 0.7, max_regions: int = 2000
) -> list[Region]:
    regions = [
        Region(
            region_id=f"{run_id}_{sheet.sheet_id}_t{i}_region",
            sheet_id=sheet.sheet_id,
            range=t.ref,
            header_row=range_boundaries(t.ref)[1] if t.headerRowCount else None,
            evidence=["Explicit Excel Table range"],
        )
        for i, t in enumerate(ws.tables.values())
    ]
    native_boxes = [range_boundaries(r.range) for r in regions]
    remaining = {
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None and not any(contains(b, cell.row, cell.column) for b in native_boxes)
    }

    def add(box, kind, confidence, evidence, features=None, header=None):
        ref = ref_for(box)
        identifier = sha256(f"{run_id}:{sheet.sheet_id}:{ref}".encode()).hexdigest()[:20]
        regions.append(
            Region(
                region_id=f"r_{identifier}",
                sheet_id=sheet.sheet_id,
                range=ref,
                region_type=kind,
                confidence=confidence,
                detected_by="deterministic_structure",
                requires_agent_review=confidence < threshold,
                evidence=evidence,
                features=features or {},
                header_row=header,
            )
        )

    # Merged text and isolated title rows are removed before table segmentation.
    for merged in ws.merged_cells.ranges:
        box = range_boundaries(str(merged))
        cell = ws.cell(box[1], box[0])
        if (cell.row, cell.column) in remaining and isinstance(cell.value, str) and cell.data_type != "f":
            kind = "narrative" if len(cell.value) > 100 else "title"
            add(box, kind, 0.94, ["Merged text block"], {"merged": True, "bold": bool(cell.font.bold)})
            remaining = {p for p in remaining if not contains(box, *p)}
    rows = defaultdict(list)
    for r, c in remaining:
        rows[r].append(c)
    for r, cols in sorted(rows.items()):
        if len(cols) == 1 and len(rows.get(r + 1, [])) >= 2:
            cell = ws.cell(r, cols[0])
            if isinstance(cell.value, str) and cell.data_type != "f":
                add((cols[0], r, cols[0], r), "title", 0.85, ["Single text row above wider block"])
                remaining.discard((r, cols[0]))
    blocks = split_blocks(remaining)
    if len(blocks) + len(regions) > max_regions:
        raise ValueError("Sheet exceeds MAX_REGIONS_PER_SHEET")
    for points in blocks:
        left, top, right, bottom = box = bounds(points)
        width, height = right - left + 1, bottom - top + 1
        cells = [ws.cell(r, c) for r, c in sorted(points)]
        density = len(points) / (width * height)
        text = [c for c in cells if isinstance(c.value, str) and c.data_type != "f"]
        formula_density = sum(c.data_type == "f" for c in cells) / len(cells)
        features = {
            "density": round(density, 3),
            "formula_density": round(formula_density, 3),
            "bold_fraction": sum(bool(c.font.bold) for c in cells) / len(cells),
            "fill_fraction": sum(c.fill.fill_type is not None for c in cells) / len(cells),
            "border_fraction": sum(
                any(s.style for s in (c.border.left, c.border.right, c.border.top, c.border.bottom))
                for c in cells
            )
            / len(cells),
            "number_formats": sorted({c.number_format for c in cells})[:10],
        }
        first = [ws.cell(top, c) for c in range(left, right + 1)]
        header = (
            width >= 2
            and all(isinstance(c.value, str) and c.data_type != "f" for c in first)
            and len({c.value for c in first}) == width
            and all(len(c.value) < 80 for c in first)
        )
        data = [ws.cell(r, c) for r, c in points if r > top]
        has_values = any(c.data_type == "f" or isinstance(c.value, (int, float)) for c in data)
        overlaps_native = any(
            not (right < b[0] or left > b[2] or bottom < b[1] or top > b[3]) for b in native_boxes
        )
        if header and height >= 2 and density >= 0.65 and not overlaps_native:
            confidence = 0.86 if has_values or height >= 3 else 0.68
            add(
                box,
                "table",
                confidence,
                ["Distinct text headers", "Contiguous rows and columns"],
                features,
                top,
            )
        elif width == 2 and formula_density > 0 and text:
            add(box, "kpi_block", 0.8, ["Label and formula layout"], features)
        elif text and len(text) == len(cells) and (width == 1 or any(len(c.value) >= 80 for c in text)):
            kind = "notes" if str(text[0].value).lower().startswith(("note", "instruction")) else "narrative"
            add(box, kind, 0.84, ["Text-only block"], features)
        else:
            add(box, "unknown", 0.5, ["No reliable header or narrative structure"], features)
    if not regions:
        add((1, 1, 1, 1), "empty", 1.0, ["No occupied cells"])
    sheet.region_count = len(regions)
    return regions
