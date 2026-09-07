from zipfile import ZipFile

from openpyxl.workbook.workbook import Workbook as ExcelWorkbook

from app.models.workbook import Sheet, Workbook


def inspect_workbook(wb: ExcelWorkbook, record: Workbook, max_cells: int) -> list[Sheet]:
    record.properties = {
        key: getattr(wb.properties, key)
        for key in ("title", "subject", "creator", "description", "created", "modified")
    }
    record.named_ranges = [
        {"name": name.name, "value": name.attr_text, "scope": scope}
        for scope, names in [
            (None, wb.defined_names),
            *[(ws.title, ws.defined_names) for ws in wb.worksheets],
        ]
        for name in names.values()
    ]
    with ZipFile(record.file_path) as archive:
        names = archive.namelist()
        record.unsupported_features = [
            label
            for prefix, label in (
                ("xl/vbaProject", "UNSUPPORTED_MACRO"),
                ("xl/externalLinks/", "EXTERNAL_WORKBOOK_REFERENCE"),
                ("xl/connections", "EXTERNAL_DATA_CONNECTION"),
                ("xl/model/", "UNSUPPORTED_DATA_MODEL"),
                ("xl/queryTables/", "UNSUPPORTED_QUERY"),
            )
            if any(n.startswith(prefix) for n in names)
        ]
    sheets = []
    for index, ws in enumerate(wb.worksheets):
        if ws.max_row * ws.max_column > max_cells:
            raise ValueError(f"Sheet {ws.title!r} exceeds MAX_SHEET_CELLS inspection limit")
        sheet = Sheet(
            sheet_id=f"{record.workbook_id}_s{index}",
            workbook_id=record.workbook_id,
            name=ws.title,
            index=index,
            visibility=ws.sheet_state,
            max_row=ws.max_row,
            max_column=ws.max_column,
            merged_ranges=[str(r) for r in ws.merged_cells.ranges],
            merged_range_count=len(ws.merged_cells.ranges),
            table_count=len(ws.tables),
            image_count=len(ws._images),
            chart_count=len(ws._charts),
        )
        for row in ws.iter_rows():
            for cell in row:
                sheet.non_empty_cells += cell.value is not None
                sheet.formula_count += cell.data_type == "f"
                if cell.comment:
                    sheet.comments.append(
                        {"cell": cell.coordinate, "text": cell.comment.text, "author": cell.comment.author}
                    )
                if cell.hyperlink:
                    sheet.hyperlinks.append(
                        {
                            "cell": cell.coordinate,
                            "target": cell.hyperlink.target,
                            "location": cell.hyperlink.location,
                        }
                    )
        sheet.native_tables = [{"name": t.name, "range": t.ref} for t in ws.tables.values()]
        sheets.append(sheet)
    record.sheet_count = len(sheets)
    for field in ("formula_count", "table_count", "image_count", "chart_count"):
        setattr(record, field, sum(getattr(s, field) for s in sheets))
    record.named_range_count = len(record.named_ranges)
    record.hidden_sheet_count = sum(s.visibility != "visible" for s in sheets)
    return sheets
