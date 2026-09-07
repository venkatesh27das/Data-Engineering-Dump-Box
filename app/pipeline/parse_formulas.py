"""Conservative reference extraction, not a formula evaluation engine."""

import re

from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils.cell import column_index_from_string
from openpyxl.workbook.workbook import Workbook as ExcelWorkbook

from app.models.assets import FormulaAsset, Reference
from app.models.workbook import Sheet

CELL = re.compile(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
COL_RANGE = re.compile(r"\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}$")
ROW_RANGE = re.compile(r"\$?[1-9][0-9]*:\$?[1-9][0-9]*$")


def address_kind(value: str) -> str | None:
    def valid_cell(part):
        match = CELL.fullmatch(part)
        return bool(match and column_index_from_string(match[1]) <= 16384 and int(match[2]) <= 1048576)

    if valid_cell(value):
        return "CELL"
    if ":" in value:
        parts = value.split(":")
        if len(parts) == 2 and all(valid_cell(p) for p in parts):
            return "RANGE"
        if COL_RANGE.fullmatch(value) and all(
            column_index_from_string(p.replace("$", "")) <= 16384 for p in parts
        ):
            return "RANGE"
        if ROW_RANGE.fullmatch(value) and all(int(p.replace("$", "")) <= 1048576 for p in parts):
            return "RANGE"
    return None


def parse_references(
    formula: str, sheet: str, wb: ExcelWorkbook
) -> tuple[list[Reference], list[str], list[str]]:
    refs, names, issues = [], [], []

    def resolve(value: str, context: str, seen: set[str]):
        target_sheet, address = context, value
        if "!" in value:
            target_sheet, address = value.rsplit("!", 1)
            if target_sheet.startswith("'") and target_sheet.endswith("'"):
                target_sheet = target_sheet[1:-1].replace("''", "'")
        if "[" in target_sheet:
            issues.append("EXTERNAL_WORKBOOK_REFERENCE")
            refs.append(
                Reference(
                    sheet=target_sheet,
                    address=address.replace("$", ""),
                    kind=address_kind(address) or "RANGE",
                    external=True,
                )
            )
            return
        if ":" in target_sheet:
            issues.append("UNSUPPORTED_3D_REFERENCE")
            return
        kind = address_kind(address)
        if kind:
            if target_sheet not in wb.sheetnames:
                issues.append("UNRESOLVED_SHEET_REFERENCE")
            refs.append(Reference(sheet=target_sheet, address=address.replace("$", "").upper(), kind=kind))
            return
        # Excel names are case insensitive; sheet-local names take precedence.
        local = wb[target_sheet].defined_names if target_sheet in wb.sheetnames else {}
        defined = next((n for n in local.values() if n.name.lower() == address.lower()), None)
        if defined is None:
            defined = next((n for n in wb.defined_names.values() if n.name.lower() == address.lower()), None)
        if defined is not None:
            key = f"{target_sheet}!{address.lower()}"
            names.append(address)
            if key in seen:
                issues.append("CYCLIC_NAMED_REFERENCE")
                return
            scan("=" + defined.attr_text.lstrip("="), target_sheet, seen | {key})
        elif "[" in address:
            issues.append("UNSUPPORTED_STRUCTURED_REFERENCE")
        else:
            issues.append("UNRESOLVED_NAMED_REFERENCE")

    def scan(expression: str, context: str, seen: set[str]):
        try:
            tokens = Tokenizer(expression).items
        except (TokenizerError, IndexError):
            issues.append("FORMULA_TOKENIZATION_FAILED")
            return
        for token in tokens:
            if (
                token.type == "FUNC"
                and token.subtype == "OPEN"
                and token.value[:-1].upper().split(".")[-1] in {"INDIRECT", "OFFSET"}
            ):
                issues.append("UNSUPPORTED_DYNAMIC_REFERENCE")
            if token.type == "OPERAND" and token.subtype == "RANGE":
                resolve(token.value, context, seen)
            if token.type == "OPERAND" and token.subtype == "ERROR" and token.value == "#REF!":
                issues.append("BROKEN_FORMULA_REFERENCE")

    scan(formula, sheet, set())
    unique = {(r.sheet, r.address, r.kind, r.external): r for r in refs}
    return list(unique.values()), sorted(set(names)), sorted(set(issues))


def extract_formulas(wb: ExcelWorkbook, value_wb: ExcelWorkbook, sheets: list[Sheet]) -> list[FormulaAsset]:
    result = []
    for sheet in sheets:
        for row in wb[sheet.name].iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", "")
                refs, names, issues = parse_references(formula, sheet.name, wb)
                if not isinstance(cell.value, str):
                    issues.append("UNSUPPORTED_ARRAY_OR_DATA_TABLE_FORMULA")
                result.append(
                    FormulaAsset(
                        formula_id=f"{sheet.sheet_id}_{cell.coordinate}",
                        workbook_id=sheet.workbook_id,
                        sheet_id=sheet.sheet_id,
                        cell=cell.coordinate,
                        formula=formula,
                        cached_value=value_wb[sheet.name][cell.coordinate].value,
                        references=refs,
                        named_references=names,
                        referenced_cells=[f"{r.sheet}!{r.address}" for r in refs if r.kind == "CELL"],
                        referenced_ranges=[f"{r.sheet}!{r.address}" for r in refs if r.kind == "RANGE"],
                        referenced_sheets=sorted({r.sheet for r in refs}),
                        external_reference=any(r.external for r in refs),
                        parse_status="partial" if issues else "parsed",
                        issues=issues,
                    )
                )
    return result
