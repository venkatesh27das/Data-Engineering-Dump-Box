import re
from pathlib import Path

import polars as pl
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from app.models.assets import TableAsset
from app.models.region import Region
from app.models.workbook import Sheet


def normalize_columns(headers: list[str]) -> list[str]:
    result, used = [], set()
    for index, header in enumerate(headers):
        base = re.sub(r"[^a-z0-9_]+", "_", header.lower()).strip("_") or f"column_{index + 1}"
        name, suffix = base, 2
        while name in used:
            name, suffix = f"{base}_{suffix}", suffix + 1
        used.add(name)
        result.append(name)
    return result


def extract_tables(
    formula_ws: Worksheet, value_ws: Worksheet, sheet: Sheet, output: Path
) -> tuple[list[TableAsset], list[Region], dict[str, pl.DataFrame]]:
    assets, regions, frames = [], [], {}
    for index, table in enumerate(formula_ws.tables.values()):
        identifier = f"{output.name}_{sheet.sheet_id}_t{index}"
        left, top, right, bottom = range_boundaries(table.ref)
        headers = [c.name for c in table.tableColumns]
        if len(headers) != right - left + 1:
            headers = [
                str(formula_ws.cell(top, c).value or f"column_{c - left + 1}") for c in range(left, right + 1)
            ]
        columns = normalize_columns(headers)
        start = top + (table.headerRowCount or 0)
        rows = list(range(start, bottom + 1))  # Includes totals and blank rows; provenance is explicit.
        warnings = []
        series = []
        for c, name in zip(range(left, right + 1), columns):
            values = [value_ws.cell(r, c).value for r in rows]
            if any(
                formula_ws.cell(r, c).data_type == "f" and value_ws.cell(r, c).value is None for r in rows
            ):
                warnings.append(f"MISSING_FORMULA_CACHE:{name}")
            try:
                series.append(pl.Series(name, values, strict=True))
            except (TypeError, ValueError, pl.exceptions.PolarsError):
                series.append(
                    pl.Series(name, [None if v is None else str(v) for v in values], dtype=pl.String)
                )
                warnings.append(f"MIXED_COLUMN_COERCED_TO_STRING:{name}")
        frame = pl.DataFrame(series)
        relative = f"tables/{identifier}.parquet"
        frame.write_parquet(output / relative)
        region = Region(region_id=f"{identifier}_region", sheet_id=sheet.sheet_id, range=table.ref)
        assets.append(
            TableAsset(
                table_id=identifier,
                workbook_id=sheet.workbook_id,
                sheet_id=sheet.sheet_id,
                region_id=region.region_id,
                table_name=table.name,
                source_range=table.ref,
                columns=columns,
                original_columns=headers,
                row_count=len(rows),
                source_rows=rows,
                parquet_path=relative,
                duckdb_table=f"table_{identifier}",
                warnings=warnings,
            )
        )
        regions.append(region)
        frames[f"table_{identifier}"] = frame
    sheet.region_count = len(regions)
    return assets, regions, frames


def extract_inferred_table(
    formula_ws: Worksheet, value_ws: Worksheet, sheet: Sheet, region: Region, output: Path
) -> TableAsset:
    left, top, right, bottom = range_boundaries(region.range)
    header_row = region.header_row
    headers = (
        [
            str(formula_ws.cell(header_row, c).value or f"column_{c - left + 1}")
            for c in range(left, right + 1)
        ]
        if header_row
        else [f"column_{c - left + 1}" for c in range(left, right + 1)]
    )
    columns = normalize_columns(headers)
    rows = list(range((header_row + 1) if header_row else top, bottom + 1))
    series, warnings = [], []
    for c, column in zip(range(left, right + 1), columns):
        values = [value_ws.cell(r, c).value for r in rows]
        if any(formula_ws.cell(r, c).data_type == "f" and value_ws.cell(r, c).value is None for r in rows):
            warnings.append(f"MISSING_FORMULA_CACHE:{column}")
        try:
            series.append(pl.Series(column, values, strict=True))
        except (TypeError, ValueError, pl.exceptions.PolarsError):
            series.append(
                pl.Series(column, [str(v) if v is not None else None for v in values], dtype=pl.String)
            )
            warnings.append(f"MIXED_COLUMN_COERCED_TO_STRING:{column}")
    identifier = f"{output.name}_{region.region_id}_table"
    relative = f"tables/{identifier}.parquet"
    pl.DataFrame(series).write_parquet(output / relative)
    if not header_row:
        warnings.append("TABLE_HEADER_UNCERTAIN")
    return TableAsset(
        table_id=identifier,
        workbook_id=sheet.workbook_id,
        sheet_id=sheet.sheet_id,
        region_id=region.region_id,
        table_name=f"Inferred {region.range}",
        source_range=region.range,
        columns=columns,
        original_columns=headers,
        row_count=len(rows),
        source_rows=rows,
        parquet_path=relative,
        duckdb_table=f"table_{identifier}",
        quality_score=region.confidence,
        warnings=warnings,
    )
