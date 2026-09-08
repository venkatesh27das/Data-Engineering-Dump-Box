"""Optional LibreOffice enrichment on temporary copies, with external links disabled."""

import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from lxml import etree


class RenderingUnavailable(RuntimeError):
    pass


def libreoffice_path() -> str | None:
    return shutil.which("soffice") or (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        if Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists()
        else None
    )


class RenderingService:
    def __init__(self, settings):
        self.settings = settings

    def convert(
        self,
        source: Path,
        destination: Path,
        output_format: str = "pdf",
        sheet_name: str | None = None,
        cell_range: str | None = None,
    ) -> Path:
        executable = libreoffice_path()
        if executable is None:
            raise RenderingUnavailable("LibreOffice is not installed")
        if output_format not in {"pdf", "xlsx"}:
            raise ValueError("Unsupported render format")
        # Reject active/external packages for optional rendering; core inspection still succeeds.
        with ZipFile(source) as archive:
            if any(
                any(
                    key in name
                    for key in (
                        "vbaProject",
                        "externalLinks/",
                        "connections.xml",
                        "queryTables/",
                        "embeddings/",
                    )
                )
                for name in archive.namelist()
            ):
                raise RenderingUnavailable("Rendering skipped for active/external workbook features")
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            for name in archive.namelist():
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    tree = etree.fromstring(archive.read(name), parser)
                    for formula in tree.findall(".//{*}f"):
                        if any(
                            token in (formula.text or "").upper()
                            for token in ("WEBSERVICE(", "RTD(", "DDE(", "HTTP:", "HTTPS:", "FILE:", "|")
                        ):
                            raise RenderingUnavailable("Rendering skipped for potentially external formulas")
        destination.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="excel-render-") as temporary:
            root = Path(temporary)
            copied = root / source.name
            shutil.copy2(source, copied)
            if sheet_name:
                import openpyxl

                wb = openpyxl.load_workbook(copied, keep_links=False)
                if sheet_name not in wb.sheetnames:
                    raise ValueError("Unknown sheet")
                for ws in wb.worksheets:
                    ws.sheet_state = "visible" if ws.title == sheet_name else "hidden"
                wb.active = wb.sheetnames.index(sheet_name)
                if cell_range:
                    wb[sheet_name].print_area = cell_range
                wb.save(copied)
                wb.close()
            profile = root / "profile"
            (profile / "user").mkdir(parents=True)
            (profile / "user/registrymodifications.xcu").write_text("""<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item></oor:items>""")
            target = root / "converted"
            target.mkdir()
            try:
                subprocess.run(
                    [
                        executable,
                        f"-env:UserInstallation={profile.as_uri()}",
                        "--headless",
                        "--norestore",
                        "--convert-to",
                        output_format,
                        "--outdir",
                        str(target),
                        str(copied),
                    ],
                    capture_output=True,
                    check=True,
                    timeout=self.settings.render_timeout_seconds,
                )
            except (OSError, subprocess.SubprocessError) as exc:
                raise RenderingUnavailable(f"LibreOffice conversion failed: {type(exc).__name__}") from exc
            result = target / f"{copied.stem}.{output_format}"
            if not result.exists():
                raise RenderingUnavailable("LibreOffice did not create an output file")
            final = destination / result.name
            if final.resolve() == source.resolve():
                raise ValueError("Cannot overwrite original workbook")
            shutil.copy2(result, final)
            return final

    def render_sheet(
        self, source: Path, destination: Path, sheet_name: str, cell_range: str | None = None
    ) -> list[Path]:
        pdf = self.convert(source, destination, sheet_name=sheet_name, cell_range=cell_range)
        import pymupdf

        images = []
        with pymupdf.open(pdf) as document:
            for index, page in enumerate(document):
                if index >= self.settings.max_render_pages:
                    break
                path = destination / f"page_{index + 1}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).save(path)
                images.append(path)
        return images

    def recalculate_copy(self, source: Path, destination: Path) -> Path:
        return self.convert(source, destination, output_format="xlsx")
