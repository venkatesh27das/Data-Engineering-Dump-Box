from pathlib import Path

import openpyxl
import polars as pl
from openpyxl.comments import Comment

from app.models.workbook import Sheet
from app.pipeline.detect_regions import detect_regions
from app.pipeline.extract_tables import extract_inferred_table
from app.pipeline.extract_text import extract_text


def sheet_record(ws):
    return Sheet(
        sheet_id="s1",
        workbook_id="w1",
        name=ws.title,
        index=0,
        visibility="visible",
        max_row=ws.max_row,
        max_column=ws.max_column,
    )


def test_separate_regions(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Regional report"
    ws.merge_cells("A1:D2")
    for r, values in enumerate([("Region", "Revenue"), ("East", 100), ("West", 200)], 4):
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
    ws["A9"] = "Management commentary with narrative content."
    ws["F4"] = "Count"
    ws["G4"] = "Value"
    ws["F5"] = 3
    ws["G5"] = 4
    ws["J12"] = 42
    ws["A5"].comment = Comment("Source note", "Tester")
    sheet = sheet_record(ws)
    regions = detect_regions(ws, sheet, "run1")
    types = {r.range: r.region_type for r in regions}
    assert types["A1:D2"] == "title"
    assert types["A4:B6"] == types["F4:G5"] == "table"
    assert types["A9:A9"] == "narrative"
    assert next(r for r in regions if r.range == "J12:J12").requires_agent_review
    (tmp_path / "tables").mkdir()
    asset = extract_inferred_table(ws, ws, sheet, next(r for r in regions if r.range == "A4:B6"), tmp_path)
    assert pl.read_parquet(tmp_path / asset.parquet_path)["revenue"].to_list() == [100, 200]
    text = extract_text(ws, sheet, regions)
    assert any(t.content_type == "comments" and "Source note" in t.content for t in text)
    assert not any(t.content == "East" for t in text)


def test_rendering_unavailable_preserves_original(settings, samples, tmp_path, monkeypatch):
    import pytest

    from app.tools.rendering_tools import RenderingService, RenderingUnavailable

    monkeypatch.setattr("app.tools.rendering_tools.libreoffice_path", lambda: None)
    source = samples["fixture_simple"]
    before = source.read_bytes()
    with pytest.raises(RenderingUnavailable, match="not installed"):
        RenderingService(settings).convert(source, tmp_path)
    assert source.read_bytes() == before


def test_rendering_runs_on_copy(settings, samples, tmp_path, monkeypatch):
    import shutil

    from app.tools.rendering_tools import RenderingService

    monkeypatch.setattr("app.tools.rendering_tools.libreoffice_path", lambda: "/fake/soffice")
    source = samples["fixture_simple"]
    original = source.read_bytes()

    def convert(command, **kwargs):
        copied = Path(command[-1])
        assert copied != source
        assert kwargs["timeout"] == settings.render_timeout_seconds
        assert "MacroSecurityLevel" in (copied.parent / "profile/user/registrymodifications.xcu").read_text()
        destination = Path(command[command.index("--outdir") + 1])
        shutil.copy2(copied, destination / copied.name)

    monkeypatch.setattr("app.tools.rendering_tools.subprocess.run", convert)
    result = RenderingService(settings).recalculate_copy(source, tmp_path / "recalculated")
    assert result.read_bytes() == original
    assert source.read_bytes() == original
