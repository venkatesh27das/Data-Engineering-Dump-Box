from io import BytesIO
from zipfile import ZipFile

import openpyxl
import polars as pl
import pytest
from openpyxl.workbook.defined_name import DefinedName

from app.graph.dependency_graph import build_graph, cell_node
from app.pipeline.extract_tables import extract_tables, normalize_columns
from app.pipeline.ingest import InvalidWorkbook, ingest
from app.pipeline.inspect_workbook import inspect_workbook
from app.pipeline.parse_formulas import extract_formulas, parse_references


def test_settings(settings):
    assert settings.upload_dir == settings.data_dir / "uploads"
    assert settings.sqlite_path == settings.data_dir / "runs.sqlite"


def test_inspection(samples, settings):
    settings.prepare()
    with samples["fixture_complex"].open("rb") as stream:
        record = ingest(stream, "../../complex.xlsx", settings)
    assert record.filename == "complex.xlsx"
    wb = openpyxl.load_workbook(record.file_path)
    sheets = inspect_workbook(wb, record, settings.max_sheet_cells)
    assert record.sheet_count == 5
    assert record.hidden_sheet_count == 2
    assert record.named_range_count == 1
    assert record.image_count == record.chart_count == 1
    assert sheets[0].comments[0]["cell"] == "A2"
    assert sheets[0].hyperlinks[0]["target"].startswith("https://example.invalid")
    assert sheets[0].merged_ranges == ["D1:F2"]
    assert [s.visibility for s in sheets][-1] == "veryHidden"
    with pytest.raises(ValueError, match="MAX_SHEET_CELLS"):
        inspect_workbook(wb, record, 1)
    wb.close()


@pytest.mark.parametrize("filename,data", [("bad.xlsm", b"bad"), ("bad.xlsx", b"bad"), ("bad.csv", b"x")])
def test_bad_upload(settings, filename, data):
    settings.prepare()
    with pytest.raises(InvalidWorkbook):
        ingest(BytesIO(data), filename, settings)
    assert not list(settings.upload_dir.rglob("*.xlsx"))


def test_size_limits(settings):
    settings.prepare()
    settings.max_upload_mb = 0.000001
    with pytest.raises(InvalidWorkbook, match="MAX_UPLOAD_MB"):
        ingest(BytesIO(b"12345"), "big.xlsx", settings)
    settings.max_upload_mb = 1
    settings.max_uncompressed_mb = 0.000001
    archive = BytesIO()
    with ZipFile(archive, "w") as z:
        z.writestr("large", "x" * 100)
    archive.seek(0)
    with pytest.raises(InvalidWorkbook, match="MAX_UNCOMPRESSED_MB"):
        ingest(archive, "big.xlsx", settings)


def test_native_extraction(samples, settings, tmp_path):
    settings.prepare()
    with samples["fixture_simple"].open("rb") as f:
        record = ingest(f, "simple.xlsx", settings)
    wb = openpyxl.load_workbook(record.file_path)
    values = openpyxl.load_workbook(record.file_path, data_only=True)
    sheets = inspect_workbook(wb, record, settings.max_sheet_cells)
    (tmp_path / "tables").mkdir()
    tables, regions, _ = extract_tables(wb["Sales"], values["Sales"], sheets[0], tmp_path)
    assert tables[0].source_rows == [2, 3]
    assert tables[0].original_columns == ["Product", "Units", "Price", "Revenue"]
    assert tables[0].warnings == ["MISSING_FORMULA_CACHE:revenue"]
    frame = pl.read_parquet(tmp_path / tables[0].parquet_path)
    assert frame["units"].to_list() == [2, 3]
    assert frame["revenue"].to_list() == [None, None]
    assert regions[0].confidence == 1
    formulas = extract_formulas(wb, values, sheets)
    assert formulas[0].formula == "=B2*C2"
    graph = build_graph(record, sheets, regions, tables, formulas)
    source = cell_node(record.workbook_id, "Summary", "B2")
    target = cell_node(record.workbook_id, "Sales", "D2:D3")
    assert graph[source][target]["relationship"] == "DEPENDS_ON"
    wb.close()
    values.close()


def test_columns():
    assert normalize_columns(["A", "a", "a_2", "", "Gross Margin %"]) == [
        "a",
        "a_2",
        "a_2_2",
        "column_4",
        "gross_margin",
    ]


@pytest.mark.parametrize(
    "formula,expected,issues",
    [
        ("=Sales!D20/Targets!$B$4", [("Sales", "D20"), ("Targets", "B4")], []),
        ("=SUM(Sales!D5:D50000)", [("Sales", "D5:D50000")], []),
        ("='Lookup Values'!C7", [("Lookup Values", "C7")], []),
        ("='O''Brien'!A1", [("O'Brien", "A1")], []),
        ('=IF(A1="Sales!A2",B1,C1)', [("Summary", "A1"), ("Summary", "B1"), ("Summary", "C1")], []),
        ('=INDIRECT("Sales!A1")', [], ["UNSUPPORTED_DYNAMIC_REFERENCE"]),
        ("=SUM(Sales!A:A)", [("Sales", "A:A")], []),
        ("=SUM(1:5)", [("Summary", "1:5")], []),
        ("='[other.xlsx]Sales'!A1", [("[other.xlsx]Sales", "A1")], ["EXTERNAL_WORKBOOK_REFERENCE"]),
        ("=SUM(Table1[Revenue])", [], ["UNSUPPORTED_STRUCTURED_REFERENCE"]),
        ("=SUM(Sales:Targets!A1)", [], ["UNSUPPORTED_3D_REFERENCE"]),
        ("=#REF!", [], ["BROKEN_FORMULA_REFERENCE"]),
    ],
)
def test_parser(formula, expected, issues):
    wb = openpyxl.Workbook()
    for name in ["Summary", "Sales", "Targets", "Lookup Values", "O'Brien"]:
        wb.create_sheet(name)
    refs, _, actual = parse_references(formula, "Summary", wb)
    assert [(r.sheet, r.address) for r in refs] == expected
    assert actual == issues


def test_names():
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.create_sheet("Sales")
    wb.defined_names.add(DefinedName("Rate", attr_text="Sales!$B$4"))
    refs, names, issues = parse_references("=rate", "Summary", wb)
    assert refs[0].address == "B4" and names == ["rate"] and not issues
    wb["Summary"].defined_names.add(DefinedName("Rate", attr_text="Sales!B5"))
    assert parse_references("=RATE", "Summary", wb)[0][0].address == "B5"
    wb.defined_names.add(DefinedName("Loop", attr_text="Loop"))
    assert "CYCLIC_NAMED_REFERENCE" in parse_references("=Loop", "Summary", wb)[2]


def test_cached_formula_value(settings, tmp_path):
    """Inject an OOXML cache to verify dual loading without running a formula engine."""
    import json

    from lxml import etree

    from app.services.run_service import RunService

    raw = BytesIO()
    wb = openpyxl.Workbook()
    wb.active["A1"] = "=2+3"
    wb.save(raw)
    raw.seek(0)
    cached = BytesIO()
    with ZipFile(raw) as source, ZipFile(cached, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = etree.fromstring(data)
                root.find(".//{*}c/{*}v").text = "5"
                data = etree.tostring(root)
            target.writestr(name, data)
    cached.seek(0)
    service = RunService(settings)
    record = service.upload(cached, "cached.xlsx")
    run = service.process(record.workbook_id)
    assert run.stage == "COMPLETED"
    formula = json.loads((settings.output_dir / run.run_id / "formulas/formulas.jsonl").read_text())
    assert formula["formula"] == "=2+3"
    assert formula["cached_value"] == 5


def test_mixed_blank_and_date_table(tmp_path):
    from datetime import datetime

    from openpyxl.worksheet.table import Table

    from app.models.workbook import Sheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Mixed", "Date"])
    ws.append([10, datetime(2024, 1, 1)])  # noqa: DTZ001 — Excel dates have no timezone
    ws.append(["ten", datetime(2024, 1, 2)])  # noqa: DTZ001 — Excel dates have no timezone
    ws.append([None, None])
    ws.add_table(Table(displayName="MixedData", ref="A1:B4"))
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    wb = openpyxl.load_workbook(path)
    sheet = Sheet(
        sheet_id="s0", workbook_id="w0", name="Sheet", index=0, visibility="visible", max_row=4, max_column=2
    )
    (tmp_path / "tables").mkdir()
    tables, _, _ = extract_tables(wb.active, wb.active, sheet, tmp_path)
    frame = pl.read_parquet(tmp_path / tables[0].parquet_path)
    assert frame.height == 3
    assert frame["mixed"].to_list() == ["10", "ten", None]
    assert frame["date"].to_list() == [datetime(2024, 1, 1), datetime(2024, 1, 2), None]  # noqa: DTZ001
    assert tables[0].warnings == ["MIXED_COLUMN_COERCED_TO_STRING:mixed"]
    wb.close()


def test_trailing_apostrophe_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "James'"
    refs, _, issues = parse_references("='James'''!A1", "James'", wb)
    assert refs[0].sheet == "James'" and not issues
