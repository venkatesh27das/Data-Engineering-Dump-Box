"""Reproducible synthetic fixtures; no Excel, LibreOffice or model required."""

import argparse
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from PIL import Image as PillowImage


def native(ws, name, ref):
    ws.add_table(Table(displayName=name, ref=ref))


def create_samples(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    paths = []

    def save(wb, name):
        path = directory / f"fixture_{name}.xlsx"
        wb.save(path)
        wb.close()
        paths.append(path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    for row in [("Product", "Units", "Price", "Revenue"), ("A", 2, 10, "=B2*C2"), ("B", 3, 20, "=B3*C3")]:
        ws.append(row)
    native(ws, "SalesData", "A1:D3")
    wb.create_sheet("Summary")["B2"] = "=SUM(Sales!D2:D3)"
    save(wb, "simple")

    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    ws["A1"] = "Regional performance"
    ws.merge_cells("A1:D2")
    for row_index, row in enumerate([("Region", "Revenue"), ("East", 100), ("West", 200)], 3):
        for column_index, value in enumerate(row, 1):
            ws.cell(row_index, column_index, value)
    native(ws, "RegionalSales", "A3:B5")
    ws["A8"] = "Synthetic management commentary for inspection."
    for row, values in enumerate([("Category", "Target"), ("Core", 300), ("Other", 50)], 11):
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    native(ws, "Targets", "A11:B13")
    save(wb, "multi_region")

    wb = Workbook()
    ws = wb.active
    ws.title = "Source"
    ws.append(["Product", "Value"])
    ws.append(["A", 100])
    ws.append(["B", 200])
    native(ws, "SourceData", "A1:B3")
    lookup = wb.create_sheet("Lookup Values")
    lookup["C7"] = 50
    summary = wb.create_sheet("Summary")
    summary["B2"] = "=SUM(Source!B2:B3)"
    summary["B3"] = "=Source!B2/'Lookup Values'!$C$7"
    save(wb, "cross_sheet")

    lookup.sheet_state = "hidden"
    wb.create_sheet("Notes")["A1"] = "Synthetic workbook for deterministic integration tests."
    archive = wb.create_sheet("Archive")
    archive.sheet_state = "veryHidden"
    ws.merge_cells("D1:F2")
    ws["D1"] = "Sales dashboard"
    ws["A2"].comment = Comment("Synthetic source row", "Fixture generator")
    ws["A2"].hyperlink = "https://example.invalid/never-fetch"
    wb.defined_names.add(DefinedName("TargetValue", attr_text="'Lookup Values'!$C$7"))
    summary["B4"] = "=Source!B3/TargetValue"
    summary["B5"] = '=INDIRECT("Source!B2")'
    chart = BarChart()
    chart.title = "Synthetic values"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "D4")
    image = BytesIO()
    PillowImage.new("RGB", (32, 32), "navy").save(image, format="PNG")
    image.seek(0)
    ws.add_image(Image(image), "J2")
    save(wb, "complex")
    return paths


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("tests/fixtures"))
    args = parser.parse_args()
    for path in create_samples(args.output):
        print(path)
