from pathlib import Path
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parents[1] / "backend" / "tests" / "fixtures" / "workbooks"


def clean_table() -> None:
    book = Workbook(); sheet = book.active; sheet.title = "Sales"
    sheet.append(["Order ID", "Region", "Revenue", "Order Date"])
    sheet.append(["ORD-1001", "East", 5800000, "2026-08-14"])
    sheet.append(["ORD-1002", "West", 4200000, "2026-08-15"])
    table = Table(displayName="SalesTable", ref="A1:D3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    sheet.add_table(table); book.save(ROOT / "01_clean_table.xlsx")


def multiple_sheets() -> None:
    book = Workbook(); customers = book.active; customers.title = "Customers"
    customers.append(["Customer ID", "Name", "Segment"]); customers.append(["CUS-1", "Aria Labs", "Enterprise"]); customers.append(["CUS-2", "Northwind", "SMB"])
    orders = book.create_sheet("Orders"); orders.append(["Order ID", "Customer ID", "Amount", "Customer Name"]); orders.append(["ORD-1", "CUS-1", 12000, '=VLOOKUP(B2,Customers!A:B,2,FALSE)']); orders.append(["ORD-2", "CUS-2", 9000, '=VLOOKUP(B3,Customers!A:B,2,FALSE)'])
    ref = book.create_sheet("Reference"); ref.sheet_state = "hidden"; ref.append(["Code", "Meaning"]); ref.append(["SMB", "Small and medium business"])
    book.save(ROOT / "02_multiple_sheets.xlsx")


def formula_model() -> None:
    book = Workbook(); inputs = book.active; inputs.title = "Inputs"; inputs.append(["Assumption", "Value"]); inputs.append(["Growth", 0.075])
    calc = book.create_sheet("Forecast"); calc.append(["Month", "Base Revenue", "Forecast Revenue"]); calc.append(["2026-08", 5800000, "=B2*(1+Inputs!B2)"]); calc["C2"].comment = Comment("Forecast uses the approved growth assumption.", "Workbook Agent")
    chart = BarChart(); chart.title = "Forecast revenue"; chart.add_data(Reference(calc, min_col=3, min_row=1, max_row=2), titles_from_data=True); chart.set_categories(Reference(calc, min_col=1, min_row=2)); calc.add_chart(chart, "E2")
    book.save(ROOT / "07_formula_model.xlsx")


def combined() -> None:
    book = Workbook(); source = book.active; source.title = "Source Data"; source.append(["Region", "Category", "Actual", "Month"])
    for row in [("East", "Electronics", 5200000, "2026-07"), ("West", "Industrial", 4100000, "2026-07")]: source.append(row)
    source.add_table(Table(displayName="SourceTable", ref="A1:D3")); source["A1"].font = Font(bold=True)
    assumptions = book.create_sheet("Assumptions"); assumptions.append(["Metric", "Value"]); assumptions.append(["Growth", 0.075])
    forecast = book.create_sheet("Regional Forecast"); forecast.append(["Region", "Category", "Forecast Month", "Forecast Revenue"]); forecast.append(["East", "Electronics", "2026-08", "='Source Data'!C2*(1+Assumptions!B2)"]); forecast.append(["West", "Industrial", "2026-08", "='Source Data'!C3*(1+Assumptions!B2)"])
    dashboard = book.create_sheet("Dashboard"); dashboard.append(["Metric", "Value"]); dashboard.append(["Total Forecast", "=SUM('Regional Forecast'!D2:D3)"]); dashboard["A5"].comment = Comment("Executive summary for planning review", "Finance")
    book.save(ROOT / "14_complex_combined.xlsx")


def advanced_layout() -> None:
    book = Workbook(); multi = book.active; multi.title = "Multi Header"
    multi.merge_cells("A1:B1"); multi["A1"] = "Customer"
    multi.merge_cells("C1:D1"); multi["C1"] = "Financial"
    multi.append(["ID", "Name", "Revenue", "Margin"])
    multi.append(["CUS-1", "Aria Labs", 120000, 0.32])
    multi.append(["CUS-2", "Northwind", 95000, 0.27])

    repeated = book.create_sheet("Repeated Blocks")
    repeated.append(["Product", "Units"]); repeated.append(["A", 10]); repeated.append(["B", 20])
    repeated.append([])
    repeated.append(["Product", "Units"]); repeated.append(["C", 30]); repeated.append(["D", 40])

    form = book.create_sheet("Employee Form")
    form.merge_cells("A1:B1"); form["A1"] = "Employee Profile"
    form.append(["Name", "Rina Das"]); form.append(["Department", "Finance"])
    form.append(["Employee ID", "EMP-104"]); form.append(["Start Date", "2025-01-15"])

    rules = book.create_sheet("Rules and Chart")
    rules.append(["Month", "Revenue", "Status"])
    rules.append(["Jan", 100, "Open"]); rules.append(["Feb", 150, "Closed"])
    rules.conditional_formatting.add(
        "B2:B3", CellIsRule(operator="greaterThan", formula=["120"])
    )
    validation = DataValidation(type="list", formula1='"Open,Closed"', allow_blank=False)
    rules.add_data_validation(validation); validation.add("C2:C20")
    chart = BarChart(); chart.title = "Monthly Revenue"
    chart.x_axis.title = "Month"; chart.y_axis.title = "Revenue"
    chart.add_data(Reference(rules, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(rules, min_col=1, min_row=2, max_row=3))
    rules.add_chart(chart, "E2")

    image_path = ROOT / "_temporary_ocr_image.png"
    picture = Image.new("RGB", (360, 120), "white")
    canvas = ImageDraw.Draw(picture)
    canvas.rectangle((8, 8, 352, 112), outline="#0b9954", width=4)
    canvas.text((24, 44), "INVOICE TOTAL 125000", fill="#122033")
    picture.save(image_path)
    rules.add_image(SpreadsheetImage(image_path), "E18")

    target = ROOT / "15_advanced_layout.xlsx"
    book.save(target)
    image_path.unlink(missing_ok=True)
    with zipfile.ZipFile(target, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "xl/externalLinks/_rels/externalLink1.xml.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/externalLinkPath" Target="file:///finance/Budget.xlsx" '
            'TargetMode="External"/></Relationships>',
        )
        archive.writestr(
            "xl/connections.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<connection id="1" name="Finance Warehouse" type="5" refreshOnLoad="1">'
            '<dbPr connection="Provider=SQL;Server=warehouse;Password=secret" '
            'command="SELECT * FROM revenue" commandType="2"/></connection></connections>',
        )
        archive.writestr(
            "xl/queryTables/queryTable1.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'name="Revenue Query" connectionId="1" refreshOnLoad="1"/>',
        )
        archive.writestr(
            "xl/pivotTables/pivotTable1.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'name="Revenue Pivot" cacheId="1"><location ref="A8:C14"/>'
            '<rowFields><field x="0"/></rowFields><dataFields>'
            '<dataField name="Sum of Revenue" fld="1" subtotal="sum"/>'
            '</dataFields></pivotTableDefinition>',
        )


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    clean_table(); multiple_sheets(); formula_model(); combined(); advanced_layout()
    print(f"Created synthetic workbook fixtures in {ROOT}")


if __name__ == "__main__":
    main()
